"""Local, read-only research tools. No network, model calls, or production writes."""
from __future__ import annotations

import argparse
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
import hashlib
import hmac
import json
import os
from pathlib import Path
import secrets
import sqlite3
from typing import Any
from zoneinfo import ZoneInfo

# Installed commands keep private data in the selected workspace, never site-packages.
ROOT = Path.cwd().resolve()
DEFAULT_SOURCE = ROOT / "data/cleaned_source.xlsx"
DEFAULT_STATE = ROOT / "research_local"
MIN_MEMBERS = 20
TOOL_VERSION = "0.2.0"
TZ = ZoneInfo("Asia/Shanghai")
# Values are fixed SQL expressions over the orders table; never user input.
GROUPS = {"all": None, "day": "order_date", "week": "week_start", "month": "month",
          "hour": "hour", "weekday": "weekday", "delivery": "delivery", "sales_source": "sales_source",
          "coupon_use": "CASE WHEN coupon_cents > 0 THEN '使用优惠券' ELSE '未使用优惠券' END",
          "wallet_use": "CASE WHEN wallet_cents > 0 THEN '使用钱包' ELSE '未使用钱包' END"}
TIME_GROUPS = {"day", "week", "month"}
PRODUCT_GROUPS = {"product": ("product",), "product_spec": ("product", "spec")}
PERIODS = {"none": None, "day": "order_date", "week": "week_start", "month": "month"}
# Candidate detail headers; the first present one is used and recorded in the manifest.
ITEM_FIELDS = {"product": ("历史商品名称", "商品名称", "商品"),
               "spec": ("历史规格名称", "规格名称", "规格", "规格属性"),
               "quantity": ("数量", "商品数量"),
               "due": ("明细应付总额", "应付金额"),
               "refund": ("明细已退金额", "已退金额"),
               "net": ("系统支付净额", "明细支付净额")}
EXPOSURE_DEFINITION = ("days_in_scope为该组时间范围与查询范围、观察期交集的日历天数；"
                       "days_with_orders为其中全店有支付观察订单的日数（不是已确认营业天数）；"
                       "per_day字段均以days_with_orders为分母。")
BUSINESS_SHEETS = ["订单数据", "商品明细", "商品价格快照", "计薪规则", "班次工时",
                   "绩效项目", "成本事件", "来源研究口径"]
LIMITATIONS = [
    "描述性统计；不能识别优惠、价格或储值的因果效应。",
    "支付观察为已支付且非明确测试，含全额退款订单；其余交易性质仍待店方确认。",
    "系统支付净额为应付减累计退款，含钱包支付；不是利润或银行流水。",
    "退款归原下单日；数据截至日可能不完整，无记录日期不能解释为零需求。",
    "只用标为会员的非空匿名ID核验分组人数下界；游客ID不计入人数门槛。",
    "金额和单量包含该分组的全部支付观察；可确认会员人数不是全体顾客数。",
]
READINESS = {
    "describe_data": {"status": "ready", "purpose": "数据清单、范围、版本和缺口"},
    "query_aggregate": {"status": "ready", "purpose": "受限订单聚合与按日/周/月时间序列，至少20位可确认会员"},
    "query_products": {"status": "ready", "purpose": "商品/规格级受限聚合，可按日/周/月；件数、明细净额和单件净额"},
    "inspect_factors": {"status": "ready", "purpose": "候选研究因素、字段依据、数据缺口和受限描述性证据；不估计影响"},
    "cohort_retention": {"status": "missing_data", "missing": [
        "用户主表及稳定匿名ID覆盖", "注册或入学时间", "覆盖目标留存期的后续订单"],
        "reason": "当前首次观察到下单不等于真实首单，更不等于新生入学。"},
    "stored_value_analysis": {"status": "missing_data", "missing": [
        "充值时间和本金", "赠送金额", "账户余额变动流水及期初/期末余额"],
        "reason": "钱包支付额只能描述支付构成，不能识别储值用户、充值或未消费余额。"},
    "ab_test_analysis": {"status": "missing_data", "missing": [
        "实验ID、分配单位和随机分组", "分配时间、资格及曝光记录", "预注册主要指标和观察窗口"],
        "reason": "优惠抵扣不能代替随机分配；用券/未用券的描述性对比可用query_aggregate(by=coupon_use)。"},
    "demand_regression": {"status": "planned", "prerequisites": [
        "确认营业日和营业时段", "价格执行事件、校历与天气", "扩大数据时间跨度并由教授确定模型"],
        "reason": "历史调价事件的前后描述可用历史数据模式的history_price_events；当前不做回归。"},
    "power_calculation": {"status": "planned", "prerequisites": [
        "主要指标、基线、最小可检测效应、显著性水平、功效及随机化单位"]},
    "make_report": {"status": "planned", "purpose": "图表、方法、代码、参数、数据版本及检验次数"},
}


class AnalysisError(ValueError):
    pass


def encoded(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False)


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def observation_calendar_days(manifest: dict) -> int | None:
    """Inclusive calendar span, not observed business days or days since cutoff."""
    try:
        start = datetime.fromisoformat(manifest["observation_start"]).date()
        end = datetime.fromisoformat(manifest["last_order_at"]).date()
        return (end - start).days + 1 if end >= start else None
    except (KeyError, ValueError, TypeError):
        return None


def private_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)
    path.chmod(0o700)


def load_key(state: Path) -> bytes:
    private_dir(state)
    path = state / "customer_hmac.key"
    if not path.exists():
        try:
            fd = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
        except FileExistsError:
            pass
        else:
            with os.fdopen(fd, "wb") as stream:
                stream.write(secrets.token_bytes(32))
    key = path.read_bytes()
    if len(key) != 32:
        raise AnalysisError("本地匿名化密钥长度异常。")
    path.chmod(0o600)
    return key


def cents(value: Any) -> int:
    if value is None or isinstance(value, bool):
        raise AnalysisError("金额缺失或类型异常，不能用零替代。")
    try:
        number = Decimal(str(value))
        if not number.is_finite():
            raise AnalysisError("金额不是有限数值。")
        integer = (number * 100).to_integral_value()
        if abs(number * 100 - integer) > Decimal("0.000001"):
            raise AnalysisError("金额精度超过分，需确认统计单位。")
        return int(integer)
    except InvalidOperation as error:
        raise AnalysisError("金额格式异常。") from error


def week_start(day: date) -> str:
    return (day - timedelta(days=day.isoweekday() - 1)).isoformat()


def order_key(identifier: Any, key: bytes) -> str | None:
    if identifier in (None, ""):
        return None
    return hmac.new(key, b"order:" + str(identifier).encode(), hashlib.sha256).hexdigest()


def normalized_order(row: dict, key: bytes) -> tuple:
    paid = row["系统已支付"]
    if type(paid) is not bool:
        raise AnalysisError("支付状态未确认。")
    created = row["下单时间"]
    if not isinstance(created, datetime):
        raise AnalysisError("下单时间缺失或不是日期时间。")
    created = created.replace(tzinfo=TZ) if created.tzinfo is None else created.astimezone(TZ)
    test = row["测试核实状态"] == "明确测试标记"
    if row["本次_支付观察"] != int(paid and not test):
        raise AnalysisError("清洗后的支付观察标记与基础字段不一致。")
    due, refund, net = [cents(row[k]) for k in ["订单应付总额", "已退金额", "系统支付净额"]]
    if net != (due - refund if paid else 0):
        raise AnalysisError("订单金额与支付状态不一致。")
    # Guests may change identifiers between devices; only member identifiers
    # provide a conservative count for the privacy threshold.
    member = None
    identifier = row["匿名顾客编号"]
    if row["顾客标识类型"] == "会员" and identifier not in (None, ""):
        member = hmac.new(key, str(identifier).encode(), hashlib.sha256).hexdigest()
    return (created.isoformat(), created.date().isoformat(), created.hour,
            created.isoweekday(), row["配送方式"] or "未记录",
            row["销售来源"] or "未记录", int(paid), int(test), member,
            net, refund, cents(row["钱包支付额"]), cents(row["优惠券减免"]),
            order_key(row.get("匿名订单编号"), key), week_start(created.date()), created.date().isoformat()[:7])


def optional_cents(value: Any) -> int | None:
    try:
        return None if value in (None, "") else cents(value)
    except AnalysisError:
        return None


def text_value(value: Any) -> str | None:
    if isinstance(value, list):
        value = "、".join(str(v) for v in value if v not in (None, ""))
    value = None if value is None else str(value).strip()
    return value or None


def normalized_items(details: list[dict], orders: list[dict], records: list[tuple]) -> tuple[list[tuple], dict]:
    """Attach item lines to their normalized parent order; the parent supplies time, payment and member key.

    Detail headers vary between exports, so fields are matched from ITEM_FIELDS and
    unmatched or unlinked lines are counted rather than guessed.
    """
    by_number = {o.get("匿名订单编号"): r for o, r in zip(orders, records) if o.get("匿名订单编号") not in (None, "")}
    by_record = {}
    for o, r in zip(orders, records):
        for field in ("源record_id", "record_id"):
            if o.get(field) not in (None, ""):
                by_record[o[field]] = r
    headers = set().union(*(d.keys() for d in details)) if details else set()
    matched = {name: next((f for f in options if f in headers), None) for name, options in ITEM_FIELDS.items()}
    items, unlinked, quantity_missing = [], 0, 0
    for detail in details:
        parent = by_number.get(detail.get("匿名订单编号"))
        links = detail.get("关联订单")
        if parent is None and isinstance(links, list) and len(links) == 1:
            parent = by_record.get(links[0])
        if parent is None or matched["product"] is None:
            unlinked += 1
            continue
        value = lambda name: detail.get(matched[name]) if matched[name] else None
        quantity = value("quantity")
        try:
            quantity = Decimal(str(quantity)) if quantity not in (None, "") and not isinstance(quantity, bool) else None
            quantity = int(quantity) if quantity is not None and quantity.is_finite() and quantity > 0 \
                and quantity == quantity.to_integral_value() else None
        except InvalidOperation:
            quantity = None
        quantity_missing += quantity is None
        items.append((parent[13], parent[0], parent[1], parent[2], parent[3], parent[14], parent[15],
                      text_value(value("product")) or "未记录", text_value(value("spec")) or "未记录",
                      quantity, optional_cents(value("due")), optional_cents(value("refund")),
                      optional_cents(value("net")), parent[6], parent[7], parent[8]))
    report = {"item_lines_in_source": len(details), "item_lines_linked": len(items),
              "item_lines_unlinked_or_unnamed": unlinked, "item_quantity_missing_lines": quantity_missing,
              "matched_item_fields": matched,
              "note": "商品明细继承父订单的时间、支付与会员标识；件数为来源记录件数，不扣除退款件数。"}
    return items, report


SCHEMA = """
CREATE TABLE orders (
    created_at TEXT NOT NULL, order_date TEXT NOT NULL, hour INTEGER NOT NULL,
    weekday INTEGER NOT NULL, delivery TEXT NOT NULL, sales_source TEXT NOT NULL,
    paid INTEGER NOT NULL, is_test INTEGER NOT NULL, member_key TEXT,
    net_cents INTEGER NOT NULL, refund_cents INTEGER NOT NULL,
    wallet_cents INTEGER NOT NULL, coupon_cents INTEGER NOT NULL,
    order_key TEXT, week_start TEXT NOT NULL, month TEXT NOT NULL
);
CREATE TABLE items (
    order_key TEXT, created_at TEXT NOT NULL, order_date TEXT NOT NULL, hour INTEGER NOT NULL,
    weekday INTEGER NOT NULL, week_start TEXT NOT NULL, month TEXT NOT NULL,
    product TEXT NOT NULL, spec TEXT NOT NULL, quantity INTEGER,
    due_cents INTEGER, refund_cents INTEGER, net_cents INTEGER,
    paid INTEGER NOT NULL, is_test INTEGER NOT NULL, member_key TEXT
);
"""
INSERT = "INSERT INTO orders VALUES (" + ",".join("?" * 16) + ")"
INSERT_ITEM = "INSERT INTO items VALUES (" + ",".join("?" * 16) + ")"


def apply_as_of(orders: list[dict], records: list[tuple], cutoff: datetime,
                as_of: str | None) -> tuple[list[dict], list[tuple], datetime, dict | None]:
    """Physically drop orders created after the as-of day so no tool can reach them."""
    if as_of is None:
        return orders, records, cutoff, None
    valid_date(as_of)
    kept = [(o, r) for o, r in zip(orders, records) if r[1] <= as_of]
    end = datetime.fromisoformat(as_of + "T23:59:59").replace(tzinfo=TZ)
    info = {"as_of": as_of, "orders_removed_after_as_of": len(orders) - len(kept),
            "rule": "建快照时已物理删除创建日期晚于as_of的订单及其商品明细；工具查不到之后的数据。"}
    return [o for o, _ in kept], [r for _, r in kept], min(cutoff, end), info


def build_snapshot(source: Path, output: Path, state: Path = DEFAULT_STATE, as_of: str | None = None) -> dict:
    import openpyxl  # Read only; never write source workbooks.

    source, output, state = source.resolve(), output.resolve(), state.resolve()
    if output.exists():
        raise AnalysisError("输出目录已存在，拒绝覆盖；请为新快照指定新目录。")
    if state == output or output in state.parents:
        raise AnalysisError("密钥与审计目录必须在快照目录之外。")
    source_hash = digest(source)
    key = load_key(state)
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        absent = set(BUSINESS_SHEETS) - set(book.sheetnames)
        if absent:
            raise AnalysisError("来源缺少预期业务表：" + "、".join(sorted(absent)))
        inventory, orders, details = [], [], []
        for name in BUSINESS_SHEETS:
            iterator = book[name].iter_rows(values_only=True)
            headers = list(next(iterator))
            if any(h is None for h in headers) or len(set(headers)) != len(headers):
                raise AnalysisError("来源表头有空值或重复：" + name)
            count = 0
            for values in iterator:
                if all(v is None for v in values):
                    continue
                count += 1
                if name == "订单数据":
                    orders.append(dict(zip(headers, values)))
                elif name == "商品明细":
                    details.append(dict(zip(headers, values)))
            inventory.append({"name": name, "row_count": count, "fields": headers})
    finally:
        book.close()
    if not orders:
        raise AnalysisError("订单表为空。")
    for field in ["record_id", "匿名订单编号"]:
        ids = [r[field] for r in orders]
        if any(v in (None, "") for v in ids) or len(set(ids)) != len(ids):
            raise AnalysisError("订单主键缺失或重复。")
    cutoffs = {r["数据截至时间"] for r in orders}
    if len(cutoffs) != 1 or not isinstance(next(iter(cutoffs)), datetime):
        raise AnalysisError("快照截至时间未统一。")
    cutoff = next(iter(cutoffs))
    cutoff = cutoff.replace(tzinfo=TZ) if cutoff.tzinfo is None else cutoff.astimezone(TZ)
    records = [normalized_order(r, key) for r in orders]
    if any(datetime.fromisoformat(r[0]) > cutoff for r in records):
        raise AnalysisError("存在晚于快照截至时间的订单。")
    orders, records, cutoff, truncation = apply_as_of(orders, records, cutoff, as_of)
    if digest(source) != source_hash:
        raise AnalysisError("读取期间来源文件发生变化，请重新导出。")
    items, item_report = normalized_items(details, orders, records)
    provenance = {"as_of_truncation": truncation} if truncation else None
    return write_analysis_snapshot(records, inventory, cutoff, source_hash, source.name, output,
                                   sum(r["匿名顾客编号"] in (None, "") for r in orders), provenance,
                                   items, item_report)


def write_analysis_snapshot(records: list[tuple], inventory: list[dict], cutoff: datetime,
                            source_hash: str, source_name: str, output: Path,
                            missing_customer_ids: int, provenance: dict | None = None,
                            items: list[tuple] = (), item_report: dict | None = None) -> dict:
    """Materialize the same read-only schema for workbook and verified Feishu inputs."""
    if not records or output.exists():
        raise AnalysisError("快照记录为空或输出目录已存在，拒绝写入。")
    output.mkdir(parents=True, mode=0o700)
    database = output / "snapshot.sqlite"
    with sqlite3.connect(database) as connection:
        connection.executescript(SCHEMA)
        connection.executemany(INSERT, records)
        connection.executemany(INSERT_ITEM, items)
    manifest = {
        "snapshot_version": f"{cutoff:%Y%m%d}_{source_hash[:12]}",
        "tool_version": TOOL_VERSION, "storage": "SQLite (v0)",
        "created_at": datetime.now(TZ).isoformat(), "timezone": "Asia/Shanghai",
        "source_file": source_name, "source_sha256": source_hash,
        "database_sha256": digest(database), "snapshot_as_of": cutoff.isoformat(),
        "observation_start": min(r[0] for r in records),
        "last_order_at": max(r[0] for r in records),
        "as_of_day_is_partial": cutoff.time().isoformat() != "23:59:59",
        "orders_in_source": len(records), "source_tables": inventory,
        "identity_coverage": {
            "missing_customer_id_orders": missing_customer_ids,
            "guest_identifiers_are_stable": False,
            "member_id_stability": "按来源会员标识暂作稳定ID；仍需业务系统确认一人一ID",
        },
        "item_coverage": item_report or {"item_lines_in_source": 0, "item_lines_linked": 0},
        "public_query_fields": {"day": "下单日期", "week": "周一起始日期", "month": "YYYY-MM",
                                "hour": "北京时间0-23时", "weekday": "星期一=1",
                                "delivery": "配送方式", "sales_source": "销售来源",
                                "coupon_use": "是否有优惠券减免", "wallet_use": "是否有钱包支付",
                                "product": "商品名称（来源原文，不合并相似名称）", "spec": "规格"},
        "privacy_min_confirmed_members": MIN_MEMBERS,
        "tools": READINESS, "limitations": LIMITATIONS,
    }
    if provenance:
        manifest["provenance"] = provenance
        if provenance.get("as_of_truncation"):
            manifest["as_of_truncation"] = provenance["as_of_truncation"]
    (output / "manifest.json").write_text(encoded(manifest), encoding="utf-8")
    database.chmod(0o400)
    (output / "manifest.json").chmod(0o400)
    return {"status": "ok", "snapshot_version": manifest["snapshot_version"],
            "snapshot_directory": str(output), "source_unchanged": True}


def read_snapshot(path: Path) -> tuple[dict, sqlite3.Connection]:
    manifest = json.loads((path / "manifest.json").read_text(encoding="utf-8"))
    database = path / "snapshot.sqlite"
    if digest(database) != manifest["database_sha256"]:
        raise AnalysisError("快照数据库哈希不一致，拒绝分析。")
    if manifest["tool_version"] != TOOL_VERSION:
        raise AnalysisError("快照与工具版本不兼容。")
    connection = sqlite3.connect(database.resolve().as_uri() + "?mode=ro&immutable=1", uri=True)
    connection.execute("PRAGMA query_only = ON")
    return manifest, connection


def valid_date(value: str | None) -> str | None:
    if value is not None:
        try:
            if date.fromisoformat(value).isoformat() != value:
                raise ValueError
        except (ValueError, TypeError):
            raise AnalysisError("日期必须是YYYY-MM-DD。") from None
    return value


def date_filters(start_date: str | None, end_date: str | None) -> tuple[list[str], list]:
    start_date, end_date = valid_date(start_date), valid_date(end_date)
    if start_date and end_date and start_date > end_date:
        raise AnalysisError("开始日期不能晚于结束日期。")
    where, parameters = ["paid = 1", "is_test = 0"], []
    for value, operator in [(start_date, ">="), (end_date, "<=")]:
        if value is not None:
            where.append(f"order_date {operator} ?")
            parameters.append(value)
    return where, parameters


def scope_dates(connection: sqlite3.Connection, start_date: str | None, end_date: str | None,
                table: str = "orders"):
    """Effective query range clipped to the observation period; table is a fixed internal name."""
    first, last = connection.execute(f"SELECT MIN(order_date), MAX(order_date) FROM {table}").fetchone()
    if first is None:
        return None, None
    start = max(start_date or first, first)
    end = min(end_date or last, last)
    return (date.fromisoformat(start), date.fromisoformat(end)) if start <= end else (None, None)


def period_bounds(by: str, label: Any, scope: tuple) -> tuple[date, date] | None:
    """Calendar dates a group covers inside the effective range (None for empty)."""
    start, end = scope
    if start is None:
        return None
    if by == "day":
        first = last = date.fromisoformat(label)
    elif by == "week":
        first = date.fromisoformat(label)
        last = first + timedelta(days=6)
    elif by == "month":
        first = date.fromisoformat(label + "-01")
        last = (first.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
    else:
        first, last = start, end
    return (first, last) if by not in TIME_GROUPS else (max(first, start), min(last, end))


def exposure(by: str, label: Any, scope: tuple, order_days: set[str]) -> dict:
    """Day counts computed by the tool so the model never has to estimate them."""
    bounds = period_bounds(by, label, scope)
    if bounds is None or bounds[0] > bounds[1]:
        return {"days_in_scope": 0, "days_with_orders": 0}
    days = [bounds[0] + timedelta(days=i) for i in range((bounds[1] - bounds[0]).days + 1)]
    if by == "weekday":
        days = [d for d in days if d.isoweekday() == int(label)]
    result = {"scope_start": bounds[0].isoformat(), "scope_end": bounds[1].isoformat(),
              "days_in_scope": len(days), "days_with_orders": sum(d.isoformat() in order_days for d in days)}
    if by in {"week", "month"}:
        full = period_bounds(by, label, (date.min, date.max))
        result["period_complete"] = bounds == full
    return result


def per_day(total: int, days: int, money: bool = False) -> float | None:
    if not days:
        return None
    value = Decimal(total) / days / (100 if money else 1)
    return float(value.quantize(Decimal("0.01")))


def query_aggregate(connection: sqlite3.Connection, by: str = "all",
                    start_date: str | None = None, end_date: str | None = None) -> dict:
    if not isinstance(by, str) or by not in GROUPS:
        raise AnalysisError("只支持" + "/".join(GROUPS) + "，不接受SQL或用户标识。")
    where, parameters = date_filters(start_date, end_date)
    # The column comes exclusively from the fixed allowlist. All user values
    # are bound parameters. There is no public raw-SQL or individual-ID input.
    column = GROUPS[by]
    grouping = column or "'全部支付观察'"
    sql = (f"SELECT {grouping}, COUNT(*), COUNT(DISTINCT member_key), "
           "SUM(net_cents), SUM(refund_cents), SUM(wallet_cents), SUM(coupon_cents), "
           "SUM(coupon_cents > 0), SUM(wallet_cents > 0) "
           "FROM orders WHERE " + " AND ".join(where))
    if column:
        sql += f" GROUP BY 1 ORDER BY 1"
    scope = scope_dates(connection, start_date, end_date)
    order_days = {d for (d,) in connection.execute(
        "SELECT DISTINCT order_date FROM orders WHERE " + " AND ".join(where), parameters)}
    rows, suppressed = [], 0
    for group, count, members, net, refund, wallet, coupon, coupon_orders, wallet_orders in \
            connection.execute(sql, parameters):
        if not count:
            continue
        # Suppressed groups contribute only to a count; no label or metric leaves this function.
        if members < MIN_MEMBERS:
            suppressed += 1
            continue
        days = exposure(by, group, scope, order_days)
        rows.append({"group": group, "paid_observed_orders": count,
                     "confirmed_members_lower_bound": members,
                     "system_net_yuan": net / 100,
                     "cumulative_refund_yuan": refund / 100,
                     "wallet_original_debit_yuan": wallet / 100,
                     "coupon_deduction_yuan": coupon / 100,
                     "orders_with_coupon": coupon_orders, "orders_with_wallet": wallet_orders,
                     "net_per_observed_order_yuan": float(
                         (Decimal(net) / Decimal(100) / count).quantize(Decimal("0.01"))),
                     **days,
                     "orders_per_day_with_orders": per_day(count, days["days_with_orders"]),
                     "net_per_day_with_orders_yuan": per_day(net, days["days_with_orders"], True)})
    notes = []
    if by == "weekday":
        notes.append("各星期的days_in_scope与days_with_orders已由工具计算；日均请直接引用per_day字段。")
    if by in {"week", "month"}:
        notes.append("period_complete为false的周期只覆盖部分日历天，不能与完整周期直接比较总量。")
    if by == "coupon_use":
        notes.append("用券与未用券为顾客自选，差异不是优惠券效果；不能替代A/B实验。")
    return {"status": "ok" if rows else "unavailable", "rows": rows,
            "privacy": {"minimum_confirmed_members": MIN_MEMBERS,
                        "notice": "仅返回满足人数门槛的分组；不返回被抑制分组的名称、金额或单量。"},
            "suppressed_group_count": suppressed, "exposure_definition": EXPOSURE_DEFINITION,
            "notes": notes,
            "reason": None if rows else "没有可发布的分组：未达到20位可确认会员或无符合条件记录。",
            "limitations": LIMITATIONS}


def query_products(connection: sqlite3.Connection, group: str = "product", period: str = "none",
                   product: str | None = None, start_date: str | None = None,
                   end_date: str | None = None, limit: int = 50) -> dict:
    if group not in PRODUCT_GROUPS or period not in PERIODS:
        raise AnalysisError("group只支持product/product_spec，period只支持none/day/week/month。")
    if product is not None and (not isinstance(product, str) or not 0 < len(product) <= 100):
        raise AnalysisError("product必须为1–100字的商品名称或null。")
    if isinstance(limit, bool) or not isinstance(limit, int) or not 1 <= limit <= 200:
        raise AnalysisError("limit必须为1–200的整数。")
    where, parameters = date_filters(start_date, end_date)
    if product is not None:
        where.append("product = ?")
        parameters.append(product)
    columns = list(PRODUCT_GROUPS[group]) + ([PERIODS[period]] if PERIODS[period] else [])
    if not connection.execute("SELECT 1 FROM items LIMIT 1").fetchone():
        return {"status": "unavailable", "rows": [], "reason": "当前快照没有可关联到订单的商品明细，不能做商品级统计。",
                "limitations": LIMITATIONS}
    sql = ("SELECT " + ", ".join(columns) + ", COUNT(*), COUNT(DISTINCT order_key), COUNT(DISTINCT member_key), "
           "SUM(quantity), SUM(quantity IS NULL), SUM(due_cents), SUM(refund_cents), SUM(net_cents), "
           "SUM(net_cents IS NULL) FROM items WHERE " + " AND ".join(where) +
           " GROUP BY " + ", ".join(columns))
    released, suppressed = [], 0
    for values in connection.execute(sql, parameters):
        labels = values[:len(columns)]
        lines, orders, members, quantity, quantity_missing, due, refund, net, net_missing = values[len(columns):]
        if members < MIN_MEMBERS:
            suppressed += 1
            continue
        row = dict(zip(["product", "spec", "period"] if group == "product_spec" else ["product", "period"], labels))
        row.update({"item_lines": lines, "orders_containing": orders, "confirmed_members_lower_bound": members,
                    "recorded_quantity": quantity, "quantity_missing_lines": quantity_missing,
                    "item_due_yuan": None if due is None else due / 100,
                    "item_refund_yuan": None if refund is None else refund / 100,
                    "item_net_yuan": None if net is None else net / 100,
                    "net_missing_lines": net_missing,
                    "net_per_unit_yuan": float((Decimal(net) / 100 / quantity).quantize(Decimal("0.01")))
                        if net is not None and quantity and not quantity_missing and not net_missing else None})
        released.append(row)
    if period == "none":
        released.sort(key=lambda r: (-(r["recorded_quantity"] or 0), r["product"]))
    else:
        released.sort(key=lambda r: (r["product"], r.get("spec") or "", r["period"]))
    return {"status": "ok" if released else "unavailable", "rows": released[:limit],
            "released_group_count": len(released), "truncated": len(released) > limit,
            "suppressed_group_count": suppressed,
            "privacy": {"minimum_confirmed_members": MIN_MEMBERS,
                        "notice": "每个商品组需至少20位可确认会员下单；不返回被抑制分组的名称或数字。"},
            "notes": ["无period时按记录件数降序排列，不代表重要性。",
                      "net_per_unit_yuan为明细系统净额除以记录件数，含优惠和退款影响，不是菜单价。",
                      "orders_containing在不同商品间重叠，不能相加得到全店订单数。"],
            "reason": None if released else "没有可发布的商品分组：未达到20位可确认会员或无符合条件记录。",
            "limitations": LIMITATIONS}


def record_audit(state: Path, entry: dict) -> None:
    private_dir(state)
    path = state / "analysis_audit.jsonl"
    fd = os.open(path, os.O_APPEND | os.O_CREAT | os.O_WRONLY, 0o600)
    with os.fdopen(fd, "a", encoding="utf-8") as stream:
        stream.write(json.dumps(entry, ensure_ascii=False, sort_keys=True, allow_nan=False) + "\n")


def run_tool(snapshot: Path, tool: str, parameters: dict | None = None,
             question: str | None = None, state: Path = DEFAULT_STATE) -> dict:
    parameters = {} if parameters is None else parameters
    started = datetime.now(TZ).isoformat()
    connection, manifest = None, {}
    try:
        if not isinstance(tool, str) or not isinstance(parameters, dict):
            raise AnalysisError("工具名必须是字符串，参数必须是对象。")
        manifest, connection = read_snapshot(snapshot)
        if tool == "describe_data":
            if parameters:
                raise AnalysisError("describe_data不接受筛选参数。")
            result = {"status": "ok", **manifest, "tools": READINESS}
        elif tool == "query_aggregate":
            if set(parameters) - {"by", "start_date", "end_date"}:
                raise AnalysisError("聚合参数不在允许清单内。")
            result = query_aggregate(connection, **parameters)
        elif tool == "query_products":
            if set(parameters) - {"group", "period", "product", "start_date", "end_date", "limit"}:
                raise AnalysisError("商品聚合参数不在允许清单内。")
            result = query_products(connection, **parameters)
        elif tool == "inspect_factors":
            if set(parameters) - {"target", "start_date", "end_date", "business_context"}:
                raise AnalysisError("因素检查参数不在允许清单内。")
            from socialdog_agent.factors import inspect_factors
            result = inspect_factors(manifest, connection, **parameters)
        elif tool in READINESS:
            result = {"status": "unavailable", "tool": tool, **READINESS[tool]}
            result["status"] = "unavailable"
        else:
            raise AnalysisError("未知分析工具。")
        result["snapshot_version"] = manifest["snapshot_version"]
        result["tool_version"] = TOOL_VERSION
        result["statistical_tests_performed"] = 0
    except (AnalysisError, OSError, KeyError, sqlite3.Error, json.JSONDecodeError) as error:
        result = {"status": "error", "error": str(error),
                  "statistical_tests_performed": 0, "tool_version": TOOL_VERSION}
    finally:
        if connection is not None:
            connection.close()
    # A failed audit write fails the call, so an unaudited result is not released.
    record_audit(state, {"started_at": started, "finished_at": datetime.now(TZ).isoformat(),
                        "question": question, "tool": tool, "parameters": parameters,
                        "snapshot_version": manifest.get("snapshot_version"),
                        "source_sha256": manifest.get("source_sha256"),
                        "tool_version": TOOL_VERSION, "result": result,
                        "result_sha256": hashlib.sha256(encoded(result).encode()).hexdigest(),
                        "statistical_tests_performed": 0})
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="SocialDog本地研究工具v0（命令行基础层）")
    parser.add_argument("--state-dir", type=Path, default=DEFAULT_STATE)
    commands = parser.add_subparsers(dest="command", required=True)
    build = commands.add_parser("build")
    build.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    build.add_argument("--output-dir", type=Path, required=True)
    build.add_argument("--as-of", help="YYYY-MM-DD；物理删除此日之后创建的订单，用于样本外测试")
    for name in ["describe", "aggregate", "cohort", "stored-value", "ab-test"]:
        sub = commands.add_parser(name)
        sub.add_argument("--snapshot", type=Path, required=True)
        sub.add_argument("--question", default=None)
        if name == "aggregate":
            sub.add_argument("--by", choices=list(GROUPS), default="all")
            sub.add_argument("--start-date")
            sub.add_argument("--end-date")
    args = parser.parse_args()
    try:
        if args.command == "build":
            result = build_snapshot(args.source, args.output_dir, args.state_dir, args.as_of)
        else:
            names = {"describe": "describe_data", "aggregate": "query_aggregate",
                     "cohort": "cohort_retention", "stored-value": "stored_value_analysis",
                     "ab-test": "ab_test_analysis"}
            parameters = {k: getattr(args, k) for k in ["by", "start_date", "end_date"]} if args.command == "aggregate" else {}
            result = run_tool(args.snapshot, names[args.command], parameters, args.question, args.state_dir)
    except (AnalysisError, OSError) as error:
        result = {"status": "error", "error": str(error)}
    print(encoded(result))
    return 0 if result["status"] == "ok" else (2 if result["status"] == "unavailable" else 1)


if __name__ == "__main__":
    raise SystemExit(main())
